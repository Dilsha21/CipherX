"""Shared helpers for building the five MFA design-document .docx files."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK = RGBColor(0x20, 0x1D, 0x18)
MUTED = RGBColor(0x5B, 0x56, 0x4A)
RULE = "DBD2BB"

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def set_cell_border(cell, color_hex=RULE, sz=4, edges=("bottom",)):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in edges:
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:color'), color_hex)
        borders.append(el)
    tcPr.append(borders)

def new_doc():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(10.5)
    style.font.color.rgb = INK
    style.paragraph_format.line_spacing = 1.28
    style.paragraph_format.space_after = Pt(6)
    sec = doc.sections[0]
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)
    return doc

def add_masthead(doc, doc_no):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("GROUP PROJECT — MULTI-FACTOR AUTHENTICATION FOR VISUALLY IMPAIRED USERS")
    r.font.size = Pt(8.5); r.font.name = 'Arial'; r.font.color.rgb = MUTED; r.bold = True
    tab = p.add_run(f"\tIndividual Design Document {doc_no} / 5")
    tab.font.size = Pt(8.5); tab.font.name = 'Arial'; tab.font.color.rgb = MUTED
    from docx.enum.text import WD_TAB_ALIGNMENT
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.2), WD_TAB_ALIGNMENT.RIGHT)
    # rule line
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(10)
    pPr = p2._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '18'); bottom.set(qn('w:color'), '201D18')
    pbdr.append(bottom)
    pPr.append(pbdr)

def add_cover(doc, tag, title, subtitle, meta, accent_hex):
    tagp = doc.add_paragraph()
    tr = tagp.add_run(tag.upper())
    tr.font.size = Pt(9); tr.font.bold = True; tr.font.name = 'Arial'
    tr.font.color.rgb = RGBColor.from_string(accent_hex)
    tagp.paragraph_format.space_after = Pt(6)

    h = doc.add_paragraph()
    hr = h.add_run(title)
    hr.font.size = Pt(22); hr.bold = True; hr.font.name = 'Arial'; hr.font.color.rgb = INK
    h.paragraph_format.space_after = Pt(6)

    sp = doc.add_paragraph()
    sr = sp.add_run(subtitle)
    sr.font.size = Pt(11); sr.italic = True; sr.font.name = 'Arial'; sr.font.color.rgb = MUTED
    sp.paragraph_format.space_after = Pt(14)

    table = doc.add_table(rows=3, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True
    labels = [k for k, v in meta]
    values = [v for k, v in meta]
    positions = [(0,0),(0,2),(1,0),(1,2),(2,0),(2,2)]
    for idx, (r, c) in enumerate(positions):
        if idx >= len(labels):
            break
        lbl_cell = table.cell(r, c)
        val_cell = table.cell(r, c+1)
        lp = lbl_cell.paragraphs[0]
        lr = lp.add_run(labels[idx].upper())
        lr.font.size = Pt(7.5); lr.bold = True; lr.font.name = 'Arial'; lr.font.color.rgb = MUTED
        vp = val_cell.paragraphs[0]
        vr = vp.add_run(values[idx])
        vr.font.size = Pt(10.5); vr.bold = True; vr.font.name = 'Arial'
        if values[idx].startswith('['):
            vr.italic = True
            vr.font.color.rgb = RGBColor.from_string(accent_hex)
        for cell in (lbl_cell, val_cell):
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            set_cell_border(cell, sz=2)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_section_heading(doc, num, title, accent_hex):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    r1 = p.add_run(f"{num}   ")
    r1.font.bold = True; r1.font.size = Pt(12); r1.font.name = 'Arial'
    r1.font.color.rgb = RGBColor.from_string(accent_hex)
    r2 = p.add_run(title.upper())
    r2.font.bold = True; r2.font.size = Pt(12); r2.font.name = 'Arial'; r2.font.color.rgb = INK
    r2.font.underline = False
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6'); bottom.set(qn('w:color'), RULE)
    pbdr.append(bottom)
    pPr.append(pbdr)

def add_para(doc, text, size=10.5):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size); r.font.name = 'Arial'; r.font.color.rgb = INK
    return p

def add_subheading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run(text)
    r.font.bold = True; r.font.size = Pt(11); r.font.name = 'Arial'; r.font.color.rgb = INK
    return p

def add_assumptions(doc, items, accent_hex):
    for aid, text in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(8)
        r1 = p.add_run(f"{aid}  ")
        r1.font.bold = True; r1.font.name = 'Arial'; r1.font.size = Pt(10.5)
        r1.font.color.rgb = RGBColor.from_string(accent_hex)
        r2 = p.add_run(text)
        r2.font.size = Pt(10.5); r2.font.name = 'Arial'; r2.font.color.rgb = INK

def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ''
        p = hdr_cells[i].paragraphs[0]
        r = p.add_run(h.upper())
        r.font.bold = True; r.font.size = Pt(8); r.font.name = 'Arial'; r.font.color.rgb = MUTED
        set_cell_border(hdr_cells[i], color_hex='201D18', sz=10, edges=("bottom",))
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            p = cells[i].paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9.5); r.font.name = 'Arial'; r.font.color.rgb = INK
            p.paragraph_format.space_after = Pt(2)
            set_cell_border(cells[i], sz=4, edges=("bottom",))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def add_diagram(doc, image_path, caption, width_cm=15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if image_path and os.path.exists(image_path):
        run = p.add_run()
        run.add_picture(image_path, width=Cm(width_cm))
    else:
        r = p.add_run("[diagram image could not be generated — see published artifact version]")
        r.italic = True; r.font.color.rgb = MUTED
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cap.add_run(caption)
    cr.italic = True; cr.font.size = Pt(9); cr.font.name = 'Arial'; cr.font.color.rgb = MUTED
    cap.paragraph_format.space_after = Pt(12)

def add_references(doc, refs):
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(ref)
        r.font.size = Pt(9.5); r.font.name = 'Arial'; r.font.color.rgb = INK

def add_footer_note(doc, doc_no):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    top = OxmlElement('w:top')
    top.set(qn('w:val'), 'single'); top.set(qn('w:sz'), '6'); top.set(qn('w:color'), RULE)
    pbdr.append(top)
    pPr.append(pbdr)
    r = p.add_run(f"Individual Contribution — Submission 1                                                                    {doc_no} / 5")
    r.font.size = Pt(8); r.font.name = 'Arial'; r.font.color.rgb = MUTED
